#!/usr/bin/env python3
"""Met en vente les t-shirts au statut « À publier » de banque-slogans.xlsx.

  python3 publish.py                 publie en APERÇU les lignes « À publier » (produit Printful + photos + page du site)
  python3 publish.py --test          fait tout à blanc : n'écrit rien chez Printful ni Cloudflare, affiche ce qui serait envoyé
  python3 publish.py publie <id>     passe ce t-shirt en vente (visible) et met le site en ligne ; « publie tout » pour tous
  python3 publish.py secrets         envoie les clés du .env à Cloudflare (sans les afficher)

Pour chaque t-shirt : 1. fichier d'impression copié à une adresse secrète du site ; 2. produit créé chez Printful
(une variante par taille, impression dans le dos) ; 3. photos du t-shirt par le générateur de mockups Printful ;
4. ajout dans catalogue.json en brouillon ; 5. site reconstruit et déployé en aperçu → lien à vérifier.
"""
import json, os, re, secrets, shutil, subprocess, sys, time
import requests
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(ROOT, "outils"))
from printful import TAILLES, ErreurPrintful, appel  # noqa: E402

SITE = os.path.join(ROOT, "site")
XLSX = os.path.join(ROOT, "banque-slogans.xlsx")
CAT = os.path.join(ROOT, "catalogue.json")
PROJET = os.getenv("CF_PROJET", "boutique")
APERCU = f"https://apercu.{PROJET}.pages.dev"
TEST = "--test" in sys.argv


def lire_json(p, defaut):
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else defaut


def ecrire_json(p, data):
    json.dump(data, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=2)


def lignes(statut):
    wb = load_workbook(XLSX)
    ws = wb["Slogans"]
    return wb, ws, [r for r in ws.iter_rows(min_row=2) if r[6].value == statut]


def fiche(id_):
    t = open(os.path.join(ROOT, "sortie", id_, "fiche-produit.md"), encoding="utf-8").read()
    titre = t.splitlines()[0].lstrip("# ").strip()
    m = re.search(r"## Description\n(.*?)\n\n(.*?)\n\n", t, re.S)
    desc = (m.group(1).strip() + "\n\n" + m.group(2).strip()) if m else ""
    return titre, desc


def deployer(branche):
    """Construit et envoie le site chez Cloudflare. branche « apercu » (lien privé) ou « main » (le vrai site)."""
    subprocess.run([sys.executable, os.path.join(ROOT, "build_site.py")] + (["--prod"] if branche == "main" else []), check=True)
    if TEST:
        print(f"  [test] wrangler pages deploy public --project-name {PROJET} --branch {branche}")
        return
    subprocess.run(["npx", "wrangler", "pages", "deploy", "public", "--project-name", PROJET, "--branch", branche,
                    "--commit-dirty=true"], cwd=SITE, check=True)


def mockups(pid_catalogue, variant_id, url_fichier, dest):
    """Photos du t-shirt porté / à plat par Printful. Renvoie la liste des fichiers enregistrés."""
    pf = appel("GET", f"/mockup-generator/printfiles/{pid_catalogue}")
    zone = next(p for p in pf["printfiles"] if p["printfile_id"] == next(
        v["placements"]["back"] for v in pf["variant_printfiles"] if v["variant_id"] == variant_id))
    w, h = zone["width"], zone["height"]
    larg = w  # notre fichier fait 4500 x 5400 (ratio 0,833) : on le cale sur la largeur, collé en haut
    haut = int(larg * 5400 / 4500)
    if haut > h:
        haut, larg = h, int(h * 4500 / 5400)
    tache = appel("POST", f"/mockup-generator/create-task/{pid_catalogue}", json={
        "variant_ids": [variant_id], "format": "jpg",
        "files": [{"placement": "back", "image_url": url_fichier,
                   "position": {"area_width": w, "area_height": h, "width": larg, "height": haut, "top": 0, "left": (w - larg) // 2}}]})
    for _ in range(30):
        time.sleep(10)
        r = appel("GET", f"/mockup-generator/task?task_key={tache['task_key']}")
        if r["status"] == "completed":
            urls = []
            for m in r["mockups"]:
                urls.append(m["mockup_url"])
                urls += [x["url"] for x in m.get("extra", [])]
            os.makedirs(dest, exist_ok=True)
            fichiers = []
            for k, u in enumerate(urls[:4]):
                p = os.path.join(dest, f"{k:02d}.jpg")
                open(p, "wb").write(requests.get(u, timeout=60).content)
                fichiers.append(os.path.relpath(p, ROOT))
            return fichiers
        if r["status"] == "failed":
            raise ErreurPrintful(f"mockups : {r.get('error')}")
    raise ErreurPrintful("mockups : Printful n'a pas répondu en 5 minutes")


def publier_un(row, cfg, cat):
    id_, style = row[9].value, row[8].value
    if not id_ or not style:
        raise ValueError("colonnes « ID design » et « Style choisi » à remplir")
    src = os.path.join(ROOT, "sortie", id_, f"impression_{style}.png")
    if not os.path.exists(src):
        raise ValueError(f"fichier introuvable : sortie/{id_}/impression_{style}.png")
    if os.path.exists(os.path.join(ROOT, "sortie", id_, "ATTENTION-BASSE-DEF.txt")):
        raise ValueError("illustration en basse définition : récupère la HD depuis Canva avant de publier")
    couleur = row[4].value or "noir"
    var = cfg["couleurs"].get(couleur) or cfg["couleurs"]["noir"]
    titre, desc = fiche(id_)
    prix = cfg["prix_vente"]
    ancien = next((c for c in cat if c["id"] == id_), {})
    code = ancien.get("code_fichier") or secrets.token_hex(8)  # 16 caractères
    dest = os.path.join(SITE, "static", "f", code)
    url = f"{APERCU}/f/{code}/{id_}.png"
    print(f"\n▶ {id_} — {titre} — {prix:.2f} € — t-shirt {couleur} ({var['printful']}), style {style}")

    # 1. fichier d'impression à une adresse introuvable sans le lien, puis mise en ligne pour que Printful le lise
    if not TEST:
        os.makedirs(dest, exist_ok=True)
        shutil.copy(src, os.path.join(dest, f"{id_}.png"))
    deployer("apercu")

    # 2. produit Printful : une variante par taille, impression dans le dos
    tailles = [t for t in TAILLES if t in var["tailles"]]
    produit = {"sync_product": {"name": titre, "external_id": id_[:32]},
               "sync_variants": [{"variant_id": var["tailles"][t], "retail_price": f"{prix:.2f}",
                                  "files": [{"type": "back", "url": url}]} for t in tailles]}
    if TEST:
        print("  [test] POST /store/products", json.dumps(produit, ensure_ascii=False)[:400], "…")
        sync = {t: None for t in tailles}
        photos = [f"sortie/{id_}/apercu_{style}.jpg"]
    else:
        r = appel("POST", "/store/products", json=produit)
        detail = appel("GET", f"/store/products/{r['id']}")
        inverse = {v: t for t, v in var["tailles"].items()}
        sync = {inverse[v["variant_id"]]: v["id"] for v in detail["sync_variants"]}
        # 3. photos
        try:
            photos = mockups(cfg["produit_id"], var["tailles"].get("M") or var["tailles"][tailles[0]], url,
                             os.path.join(SITE, "img", id_))
        except (ErreurPrintful, StopIteration, KeyError) as e:
            print(f"  ⚠ photos Printful indisponibles ({e}) : j'utilise l'aperçu maison en attendant")
            photos = [f"sortie/{id_}/apercu_{style}.jpg"]

    # 4. catalogue (brouillon)
    entree = {"id": id_, "titre": titre, "description": desc, "niche": niche(id_, row[1].value), "couleur": couleur,
              "prix": prix, "photos": photos, "visible": ancien.get("visible", False), "style": style, "code_fichier": code,
              "tailles": [{"taille": t, "sync_variant_id": sync[t]} for t in tailles]}
    cat[:] = [c for c in cat if c["id"] != id_] + [entree]
    if not TEST:
        ecrire_json(CAT, cat)
    return entree


def niche(id_, niche_excel):
    sys.path.insert(0, ROOT)
    from build_site import niche_de
    n = (niche_excel or "").lower()
    if "corse" in n or "région" in n:
        return "Régions"
    if "apéro" in n:
        return "Apéro"
    return niche_de(id_)


def publier_apercu():
    cfg = lire_json(os.path.join(ROOT, "config_printful.json"), None)
    if not cfg:
        sys.exit("Choisis d'abord le t-shirt de base : python3 choisir_tshirt.py")
    wb, ws, a_faire = lignes("À publier")
    if not a_faire:
        print("Aucune ligne « À publier » dans banque-slogans.xlsx.")
        return
    cat = lire_json(CAT, [])
    faits, erreurs = [], []
    for row in a_faire:
        try:
            faits.append(publier_un(row, cfg, cat))
            if not TEST:
                row[6].value = "En aperçu"  # évite de recréer le produit si on relance
                row[7].value = f"Aperçu : {APERCU}/t-shirt/{row[9].value}/"
        except (ErreurPrintful, ValueError, requests.RequestException, subprocess.CalledProcessError) as e:
            erreurs.append((row[9].value or row[0].value, str(e)))
            print(f"  ✖ arrêté pour ce t-shirt : {e}")
    if faits:
        deployer("apercu")
    if not TEST:
        wb.save(XLSX)
    print("\nRésumé :")
    for e in faits:
        print(f"  ✔ {e['titre']} — {e['prix']:.2f} € — {APERCU}/t-shirt/{e['id']}/")
    for i, m in erreurs:
        print(f"  ✖ {i} : {m}")
    if faits:
        print("Vérifie les aperçus, puis dis « publie » (python3 publish.py publie tout).")


def mettre_en_vente(qui):
    cat = lire_json(CAT, [])
    cibles = [c for c in cat if qui == "tout" and not c.get("visible") or c["id"] == qui]
    if not cibles:
        sys.exit(f"Rien à publier pour « {qui} » (déjà en ligne, ou pas encore passé par l'aperçu).")
    for c in cibles:
        c["visible"] = True
    if not TEST:
        ecrire_json(CAT, cat)
    deployer("main")
    wb, ws, _ = lignes("En aperçu")
    ids = {c["id"] for c in cibles}
    for r in ws.iter_rows(min_row=2):
        if r[9].value in ids:
            r[6].value = "En ligne"
    if not TEST:
        wb.save(XLSX)
    site = os.getenv("SITE_URL", "").rstrip("/")
    print("En vente ! Liens à mettre sous les vidéos :")
    for c in cibles:
        print(f"  {c['titre']} → {site}/t-shirt/{c['id']}/")


def envoyer_secrets():
    from dotenv import dotenv_values
    v = dotenv_values(os.path.join(ROOT, ".env"))
    for k in ("STRIPE_SECRET_KEY", "STRIPE_WEBHOOK_SECRET", "PRINTFUL_TOKEN", "PRINTFUL_STORE_ID", "SITE_URL",
              "PRINTFUL_AUTO_CONFIRM", "RESEND_API_KEY", "EMAIL_MAXIME"):
        if v.get(k):
            subprocess.run(["npx", "wrangler", "pages", "secret", "put", k, "--project-name", PROJET],
                           input=v[k].encode(), cwd=SITE, check=True, stdout=subprocess.DEVNULL)
            print(f"  ✔ {k} envoyé à Cloudflare")


if __name__ == "__main__":
    a = [x for x in sys.argv[1:] if x != "--test"]
    if a[:1] == ["publie"]:
        mettre_en_vente(a[1] if len(a) > 1 else "tout")
    elif a[:1] == ["secrets"]:
        envoyer_secrets()
    else:
        publier_apercu()
